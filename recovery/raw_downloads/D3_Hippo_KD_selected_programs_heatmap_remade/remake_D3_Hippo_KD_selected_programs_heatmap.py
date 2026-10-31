from pathlib import Path
from openpyxl import load_workbook
import numpy as np
import pandas as pd
import matplotlib as mpl
mpl.use('Agg')
import matplotlib.pyplot as plt

# Illustrator-editable vector output
mpl.rcParams.update({
    'font.family': 'DejaVu Sans',
    'font.size': 10,
    'axes.titlesize': 18,
    'axes.labelsize': 12,
    'xtick.labelsize': 13,
    'ytick.labelsize': 14,
    'pdf.fonttype': 42,
    'ps.fonttype': 42,
    'svg.fonttype': 'none',
})

XLSX = Path('/mnt/data/cNMF_50_0_2_032526_6co.xlsx')
OUT = Path('/mnt/data/D3_Hippo_KD_selected_programs_heatmap_remade')
OUT.mkdir(exist_ok=True)

SHEET = 'Perturbation Association 1'
ROWS = ['NF2', 'LATS2', 'AMOTL2', 'TAOK1']
PROGRAMS = ['P2', 'P3', 'P8', 'P11', 'P33', 'P34']
PROGRAM_LABELS = {
    'P2': 'P2\nArtery\nendothelium 1',
    'P3': 'P3\nEndothelial\nfunction 1',
    'P8': 'P8\nArtery\nendothelium 2',
    'P11': 'P11\nEndothelial\nfunction 2',
    'P33': 'P33\nAnterior-like\nmesoderm',
    'P34': 'P34\nMesenchyme,\nYAP/TAZ-induced',
}

VALUE_COL = 'approx_log2FC'
FDR_COL = 'adj_pval'
SAMPLE = 'D3'

# Read workbook with openpyxl to avoid altering source file
wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb[SHEET]
headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
records = []
for r in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(headers, r))
    target = str(d.get('target_name'))
    program = f"P{int(float(d.get('program_name')))}" if d.get('program_name') is not None else None
    sample = str(d.get('Sample'))
    if target in ROWS and program in PROGRAMS and sample == SAMPLE:
        records.append({
            'target_name': target,
            'program': program,
            'program_label': PROGRAM_LABELS[program].replace('\n', ' '),
            'approx_log2FC': float(d.get('approx_log2FC')) if d.get('approx_log2FC') is not None else np.nan,
            'beta_hat': float(d.get('beta_hat')) if d.get('beta_hat') is not None else np.nan,
            'p_value': float(d.get('p-value')) if d.get('p-value') is not None else np.nan,
            'adj_pval': float(d.get('adj_pval')) if d.get('adj_pval') is not None else np.nan,
            'specificity_scores': float(d.get('specificity_scores')) if d.get('specificity_scores') is not None else np.nan,
            'Sample': sample,
        })

source = pd.DataFrame(records)
if source.shape[0] != len(ROWS) * len(PROGRAMS):
    raise ValueError(f'Expected {len(ROWS)*len(PROGRAMS)} rows but found {source.shape[0]}')
source['target_name'] = pd.Categorical(source['target_name'], categories=ROWS, ordered=True)
source['program'] = pd.Categorical(source['program'], categories=PROGRAMS, ordered=True)
source = source.sort_values(['target_name', 'program']).reset_index(drop=True)
source.to_csv(OUT / 'D3_Hippo_KD_selected_programs_source.tsv', sep='\t', index=False)

mat = source.pivot(index='target_name', columns='program', values=VALUE_COL).loc[ROWS, PROGRAMS]
qmat = source.pivot(index='target_name', columns='program', values=FDR_COL).loc[ROWS, PROGRAMS]
pmat = source.pivot(index='target_name', columns='program', values='p_value').loc[ROWS, PROGRAMS]
mat.to_csv(OUT / 'D3_Hippo_KD_selected_programs_approx_log2FC_matrix.tsv', sep='\t')
qmat.to_csv(OUT / 'D3_Hippo_KD_selected_programs_adj_pval_matrix.tsv', sep='\t')
pmat.to_csv(OUT / 'D3_Hippo_KD_selected_programs_pvalue_matrix.tsv', sep='\t')

def star(q):
    if np.isnan(q):
        return ''
    if q < 0.001:
        return '***'
    if q < 0.01:
        return '**'
    if q < 0.05:
        return '*'
    return ''

def plot_heatmap(show_values=False, stem_suffix='stars_only'):
    vals = mat.to_numpy(dtype=float)
    vmax = max(6, float(np.nanmax(np.abs(vals))))
    vmin = -vmax
    n_rows, n_cols = vals.shape

    # square cells: axes width/height ratio follows n_cols/n_rows
    cell = 0.72
    fig_w = 1.9 + cell * n_cols + 1.15
    fig_h = 1.0 + cell * n_rows + 1.35
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    mesh = ax.pcolormesh(
        np.arange(n_cols + 1),
        np.arange(n_rows + 1),
        vals,
        cmap='RdBu_r',
        vmin=vmin,
        vmax=vmax,
        edgecolors='white',
        linewidth=1.5,
        shading='flat',
    )
    ax.set_aspect('equal')
    ax.set_xlim(0, n_cols)
    ax.set_ylim(n_rows, 0)
    ax.set_xticks(np.arange(n_cols) + 0.5)
    ax.set_yticks(np.arange(n_rows) + 0.5)
    ax.set_xticklabels([PROGRAM_LABELS[p] for p in PROGRAMS], rotation=42, ha='right', va='top')
    ax.set_yticklabels(ROWS)
    ax.tick_params(length=0)
    for sp in ax.spines.values():
        sp.set_visible(False)
    ax.set_title('D3 cNMF program effects of Hippo gene KD', pad=12, fontweight='bold')

    for i, target in enumerate(ROWS):
        for j, program in enumerate(PROGRAMS):
            v = mat.loc[target, program]
            q = qmat.loc[target, program]
            s = star(q)
            if show_values:
                txt = f'{v:.2f}{s}' if s else f'{v:.2f}'
            else:
                txt = s
            if txt:
                color = 'white' if abs(v) > 0.55 * vmax else 'black'
                ax.text(j + 0.5, i + 0.5, txt, ha='center', va='center', fontsize=13, color=color, fontweight='bold' if s and not show_values else 'normal')

    cbar = fig.colorbar(mesh, ax=ax, pad=0.03, fraction=0.045)
    cbar.set_label('D3 approx. log2FC', fontsize=14)
    cbar.ax.tick_params(labelsize=11)
    # Star legend, as separate editable text
    fig.text(0.82, 0.18, '* FDR < 0.05\n** FDR < 0.01\n*** FDR < 0.001', ha='left', va='bottom', fontsize=12)

    fig.subplots_adjust(left=0.17, bottom=0.28, right=0.79, top=0.86)
    stem = f'D3_Hippo_KD_selected_programs_P2_P3_P8_P11_P33_P34_{stem_suffix}'
    for ext in ['svg', 'pdf', 'png']:
        fig.savefig(OUT / f'{stem}.{ext}', bbox_inches='tight', dpi=300 if ext == 'png' else None)
    plt.close(fig)

plot_heatmap(show_values=False, stem_suffix='stars_only')
plot_heatmap(show_values=True, stem_suffix='values_and_stars')

(OUT / 'README.md').write_text('''# D3 Hippo KD selected program heatmap\n\nInput workbook: cNMF_50_0_2_032526_6co.xlsx\nSheet: Perturbation Association 1\nSample: D3\nRows: NF2, LATS2, AMOTL2, TAOK1\nColumns: P2, P3, P8, P11, P33, P34\nValue: approx_log2FC\nStars: adj_pval thresholds (* < 0.05, ** < 0.01, *** < 0.001)\n\nTwo SVG/PDF/PNG versions are provided: stars-only and values plus stars. SVG text remains editable in Illustrator.\n''')
print('wrote', OUT)
print(source[['target_name','program','approx_log2FC','adj_pval']].to_string(index=False))
