import re, shutil, glob, zipfile
from pathlib import Path
import numpy as np, pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.gridspec import GridSpec
from matplotlib.patches import FancyBboxPatch
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

plt.rcParams.update({'font.family':'DejaVu Sans','svg.fonttype':'none','pdf.fonttype':42,'ps.fonttype':42})
BASE=Path('/mnt/data')
OUT=BASE/'figure1_revised'
COL_REQUIRED='#2B6CB0'; COL_RESTRICT='#D55E00'

def export_fig(fig, stem, dpi=350):
    stem=Path(stem)
    fig.savefig(str(stem)+'.svg', bbox_inches='tight')
    fig.savefig(str(stem)+'.pdf', bbox_inches='tight')
    fig.savefig(str(stem)+'.png', dpi=dpi, bbox_inches='tight')
    plt.close(fig)

gene=pd.read_csv(OUT/'cleaned_data'/'genomewide_results_unique_gene_level.tsv', sep='\t')
# Harmonize column names from earlier builds if needed.
if 'screen_effect' in gene.columns and 'mageck_lfc' not in gene.columns:
    gene['mageck_lfc'] = gene['screen_effect']
if 'best_mageck_fdr' in gene.columns and 'directional_fdr' not in gene.columns:
    gene['directional_fdr'] = gene['best_mageck_fdr']
if 'hit_support' in gene.columns and 'method_support' not in gene.columns:
    gene['method_support'] = gene['hit_support']
if 'screen_direction' in gene.columns and 'direction' not in gene.columns:
    gene['direction'] = gene['screen_direction'].map({'Required for EC differentiation':'required_for_EC_differentiation','Restricts EC differentiation':'restricts_EC_differentiation'}).fillna(gene['screen_direction'])
anno=pd.read_csv('/mnt/data/figure1_annotation_literature_draft/screen_regulators_244_literature_annotation_draft.tsv', sep='\t')
# -------- Panel 1b revised --------
d=gene.copy()
if 'directional_fdr' not in d.columns:
    d['directional_fdr'] = np.where(d['mageck_lfc'] < 0, d['neg|fdr'], d['pos|fdr'])
d['neglog10']=-np.log10(d['directional_fdr'].clip(lower=1e-300))
d['neglog10_capped']=d['neglog10'].clip(upper=4.35)
fig,ax=plt.subplots(figsize=(6.6,4.8))
nh=d[~d['union_hit']]
ax.scatter(nh['mageck_lfc'], nh['neglog10_capped'], s=4, c='#D8D8D8', alpha=0.22, linewidth=0, rasterized=True)
for direction,color,label in [('required_for_EC_differentiation',COL_REQUIRED,'Required for EC differentiation'),('restricts_EC_differentiation',COL_RESTRICT,'Restricts EC differentiation')]:
    sub=d[d['union_hit'] & (d['direction']==direction)]
    ax.scatter(sub['mageck_lfc'], sub['neglog10_capped'], s=18, c=color, alpha=0.86, edgecolor='white', linewidth=0.25, label=label)
both=d[d['union_hit'] & (d['method_support']=='both')]
ax.scatter(both['mageck_lfc'], both['neglog10_capped'], s=31, facecolors='none', edgecolors='black', linewidths=0.45, label='Both methods')
# prioritized labels only; manual offsets
labels=['CDH5','KDR','ETV2','SMAD1','SMAD5','ACVR2A','PLCG1','BMPR1A','AMOTL2']
offsets={'CDH5':(-26,16),'KDR':(12,-6),'ETV2':(16,12),'SMAD1':(-30,4),'SMAD5':(-20,-14),'ACVR2A':(10,20),'PLCG1':(12,6),'BMPR1A':(18,-18),'AMOTL2':(10,0)}
for g in labels:
    sub=d[d['normalized_gene_symbol']==g]
    if len(sub) and sub.iloc[0]['union_hit']:
        r=sub.iloc[0]
        ax.annotate(g,(r['mageck_lfc'],r['neglog10_capped']),xytext=offsets.get(g,(5,5)),textcoords='offset points',fontsize=7.2,arrowprops=dict(arrowstyle='-',lw=0.35,color='#777'),ha='center')
ax.axvline(0,color='#999',ls='--',lw=0.7)
ax.axhline(-np.log10(0.05),color='#999',ls=':',lw=0.7)
ax.set_title('Genome-wide screen result',fontsize=10.7,fontweight='bold')
ax.set_xlabel('MaGECK log2 fold change (CDH5+ vs CDH5-)',fontsize=9)
ax.set_ylabel('-log10(direction-matched FDR)',fontsize=9)
ax.set_xlim(-3.7,2.7); ax.set_ylim(-0.1,4.55)
ax.text(0.02,0.97,'244 unique screen-derived regulators\nY-axis capped at 4.35 for display',transform=ax.transAxes,ha='left',va='top',fontsize=7.4)
ax.tick_params(labelsize=8)
ax.spines[['top','right']].set_visible(False)
ax.legend(loc='upper right',fontsize=6.5,frameon=False,markerscale=0.75)
export_fig(fig, OUT/'individual_panels'/'Fig1b_selected_main_panel')
# overwrite candidate selected too
for ext in ['svg','pdf','png']:
    shutil.copyfile(OUT/'individual_panels'/f'Fig1b_selected_main_panel.{ext}', OUT/'candidate_panels'/f'Fig1b_candidate_B_MaGECK_plot.{ext}')

# -------- Panel 1d revised --------
selected_cats=[('VEGF / endothelial','VEGF/endothelial signaling'),('TGF-beta / BMP','TGF-beta/BMP signaling'),('Hippo / YAP','Hippo/YAP signaling'),('SWI/SNF chromatin','Chromatin modification/remodeling'),('Mediator / transcription','Mediator/general transcription machinery')]
plot_anno=anno.merge(gene[['normalized_gene_symbol','mageck_lfc','union_hit','direction','method_support']],left_on='gene',right_on='normalized_gene_symbol',how='left')
rows=[]
key_label=set(['CDH5','KDR','ETV2','PLCG1','SMAD1','SMAD4','SMAD5','ACVR2A','BMPR1A','NF2','LATS2','AMOTL2','SAV1','TAOK1','SMARCA4','SMARCC1','MED12'])
for display,sec in selected_cats:
    sub=plot_anno[plot_anno['secondary_category']==sec].copy()
    sub=sub.sort_values('mageck_lfc', key=lambda s:s.abs(), ascending=False).head(10)
    for i,(_,r) in enumerate(sub.iterrows()):
        label=(r['gene'] in key_label) or i<2
        rows.append({'pathway_display_name':display,'secondary_category':sec,'member_gene':r['gene'],'screen_effect':r['mageck_lfc'],'gene_is_final_hit':r['union_hit'],'direction':r['direction'],'display_label':label})
comp=pd.DataFrame(rows)
comp.to_csv(OUT/'source_data'/'Fig1d_component_effects.tsv', sep='\t', index=False)
comp.to_csv(OUT/'source_data'/'Fig1d_pathway_membership.tsv', sep='\t', index=False)
fig,axes=plt.subplots(len(selected_cats),1,figsize=(6.8,4.9),sharex=True,gridspec_kw={'hspace':0.42})
for ax,(display,sec) in zip(axes,selected_cats):
    sub=comp[comp['pathway_display_name']==display].sort_values('screen_effect')
    y=np.zeros(len(sub))
    ax.axhline(0,color='#D0D0D0',lw=0.6)
    ax.axvline(0,color='#A0A0A0',lw=0.7,ls='--')
    colors=np.where(sub['screen_effect']<0,COL_REQUIRED,COL_RESTRICT)
    ax.scatter(sub['screen_effect'], y, s=40, c=colors, edgecolor='white', lw=0.35, zorder=3)
    # jitter slightly vertical for collisions
    for j,(_,r) in enumerate(sub.iterrows()):
        if r['display_label']:
            dy=0.08 if j%2==0 else -0.08
            ha='left' if r['screen_effect']<0 else 'right'
            dx=0.03 if r['screen_effect']<0 else -0.03
            ax.text(r['screen_effect']+dx, dy, r['member_gene'], fontsize=6.7, ha=ha, va='center')
    ax.set_ylim(-0.24,0.24); ax.set_yticks([])
    ax.text(-3.55,0.17,display,fontsize=8.2,fontweight='bold',ha='left',va='center')
    ax.spines[['top','right','left','bottom']].set_visible(False)
    ax.tick_params(axis='x',labelsize=7,length=2)
axes[-1].spines['bottom'].set_visible(True)
axes[-1].set_xlabel('MaGECK log2 fold change (CDH5+ vs CDH5-)',fontsize=8.5)
fig.suptitle('Screen hits converge on selected pathway components',fontsize=11,fontweight='bold',y=0.99)
# legend text
fig.text(0.03,0.025,'Blue = required for EC differentiation; orange = restricts EC differentiation. Points are screen-derived hit components from literature-informed annotation groups.',fontsize=6.6,color='#555')
export_fig(fig, OUT/'individual_panels'/'Fig1d_selected_main_panel')
# also alias named
for ext in ['svg','pdf','png']:
    shutil.copyfile(OUT/'individual_panels'/f'Fig1d_selected_main_panel.{ext}', OUT/'individual_panels'/f'Fig1d_selected_pathway_component_strips.{ext}')

# -------- Reassemble --------
def load_img(path): return Image.open(path).convert('RGB')
imgs={k:load_img(OUT/'individual_panels'/v) for k,v in {'a':'Fig1a_screen_workflow.png','b':'Fig1b_selected_main_panel.png','d':'Fig1d_selected_main_panel.png','e':'Fig1e_selected_main_panel.png','c':'Fig1c_screen_vs_abundance_placeholder.png'}.items()}
fig=plt.figure(figsize=(11.5,8.0))
gs=GridSpec(2,2,figure=fig,width_ratios=[1,1.25],height_ratios=[0.92,1.08],wspace=0.08,hspace=0.13)
for key,pos,letter in [('a',gs[0,0],'a'),('b',gs[0,1],'b'),('d',gs[1,0],'c'),('e',gs[1,1],'d')]:
    ax=fig.add_subplot(pos); ax.imshow(imgs[key]); ax.axis('off'); ax.text(-0.02,1.02,letter,transform=ax.transAxes,fontsize=16,fontweight='bold',va='bottom')
fig.savefig(OUT/'assembled_figures'/'Figure1_revised_fallback.png',dpi=350,bbox_inches='tight')
fig.savefig(OUT/'assembled_figures'/'Figure1_revised_fallback.pdf',bbox_inches='tight')
fig.savefig(OUT/'assembled_figures'/'Figure1_revised_fallback.svg',bbox_inches='tight')
plt.close(fig)
# 5-panel internal
fig=plt.figure(figsize=(12.2,8.2))
gs=GridSpec(2,3,figure=fig,width_ratios=[1.05,1.25,1.1],height_ratios=[0.95,1.05],wspace=0.08,hspace=0.13)
for key,pos,letter in [('a',gs[0,0],'a'),('b',gs[0,1:],'b'),('c',gs[1,0],'c'),('d',gs[1,1],'d'),('e',gs[1,2],'e')]:
    ax=fig.add_subplot(pos); ax.imshow(imgs[key]); ax.axis('off'); ax.text(-0.02,1.02,letter,transform=ax.transAxes,fontsize=16,fontweight='bold',va='bottom')
fig.savefig(OUT/'assembled_figures'/'Figure1_revised_internal_placeholder.png',dpi=350,bbox_inches='tight')
fig.savefig(OUT/'assembled_figures'/'Figure1_revised_internal_placeholder.pdf',bbox_inches='tight')
fig.savefig(OUT/'assembled_figures'/'Figure1_revised_internal_placeholder.svg',bbox_inches='tight')
plt.close(fig)
Image.open(OUT/'assembled_figures'/'Figure1_revised_fallback.png').save(OUT/'assembled_figures'/'Figure1_revised_fallback_600dpi.tiff',dpi=(600,600))
# pptx
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.shapes.add_picture(str(OUT/'assembled_figures'/'Figure1_revised_fallback.png'),Inches(0.2),Inches(0.2),width=Inches(12.9))
prs.save(OUT/'assembled_figures'/'Figure1_revised_fallback.pptx')
# recreate zips
for name, pattern in [('individual_panel_zip','individual_panels/*'),('candidate_panel_zip','candidate_panels/*'),('source_data_zip','source_data/*')]:
    with zipfile.ZipFile(OUT/f'{name}.zip','w',zipfile.ZIP_DEFLATED) as z:
        for fp in glob.glob(str(OUT/pattern)):
            z.write(fp, arcname=Path(fp).relative_to(OUT))
with zipfile.ZipFile(BASE/'figure1_revised_full_directory.zip','w',zipfile.ZIP_DEFLATED) as z:
    for fp in OUT.rglob('*'):
        if fp.is_file(): z.write(fp, arcname=fp.relative_to(OUT.parent))
print('revised visuals done')
