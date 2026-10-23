import os, re, json, zipfile, shutil, time
from pathlib import Path
import numpy as np, pandas as pd
import matplotlib as mpl
mpl.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Circle, Rectangle
from matplotlib.colors import Normalize, TwoSlopeNorm
from scipy.stats import pearsonr, spearmanr
from openpyxl import load_workbook
try:
    from adjustText import adjust_text
except Exception:
    adjust_text=None
mpl.rcParams['pdf.fonttype']=42; mpl.rcParams['ps.fonttype']=42; mpl.rcParams['svg.fonttype']='none'; mpl.rcParams['font.family']='DejaVu Sans'
BASE=Path('/mnt/data'); OUT=BASE/'figure2_build_simple'
if OUT.exists(): shutil.rmtree(OUT)
for d in ['data_audit','cleaned_data','source_data','scripts','candidate_panels','individual_panels','assembled_figure','legends']:
    (OUT/d).mkdir(parents=True,exist_ok=True)
CNMF=BASE/'cNMF_50_0_2_032526_6co(1).xlsx'
CONC=BASE/'u_test_vs_D3_EC_fraction_AI_editable/u_test_vs_D3_EC_fraction_merged_source_AI_editable.csv'
print('loading workbook', flush=True)
wb=load_workbook(CNMF, read_only=True, data_only=True)
def read_sheet(name):
    print('read',name,flush=True)
    rows=list(wb[name].iter_rows(values_only=True))
    header=[str(h) if h is not None else f'unnamed_{i}' for i,h in enumerate(rows[0])]
    data=[r for r in rows[1:] if any(x is not None for x in r)]
    return pd.DataFrame(data,columns=header)
summary=read_sheet('Summary'); loadings=read_sheet('Program Loadings'); targets=read_sheet('Targets Summary'); assoc=read_sheet('Perturbation Association 1')
summary['program_name']=pd.to_numeric(summary['program_name'],errors='coerce').astype(int)
loadings['Program']=pd.to_numeric(loadings['Program'],errors='coerce').astype('Int64'); loadings['Rank']=pd.to_numeric(loadings['Rank'],errors='coerce')
assoc['program_name']=pd.to_numeric(assoc['program_name'],errors='coerce').astype(int); assoc['approx_log2FC']=pd.to_numeric(assoc['approx_log2FC'],errors='coerce'); assoc['adj_pval']=pd.to_numeric(assoc['adj_pval'],errors='coerce'); assoc['p-value']=pd.to_numeric(assoc['p-value'],errors='coerce'); assoc['Sample']=assoc['Sample'].astype(str)
print('data loaded', summary.shape, loadings.shape, targets.shape, assoc.shape, flush=True)
# metadata
prog_labels={7:('D0','P7','pluripotency'),22:('D0','P22','TDGF1/pluripotency'),13:('D1','P13','posterior primitive streak'),30:('D1','P30','anterior primitive streak'),15:('D2','P15','Pre-EC / pre-artery'),17:('D2','P17','Lat-Meso / epithelial'),19:('D2','P19','posterior mesoderm'),23:('D2','P23','EC induction'),2:('D3','P2','EC endpoint'),3:('D3','P3','EC endpoint'),8:('D3','P8','vascular EC'),11:('D3','P11','EC activation'),27:('D3','P27','mesenchymal alternative'),34:('D3','P34','contractile / YAP-like')}
selected=[7,22,13,30,15,17,19,23,2,3,8,11,27,34]
POS='#D95F02'; NEG='#1F78B4'; PURPLE='#7B3294'; GREY='#C9C9C9'; DAY={'D0':'#555555','D1':'#3A7D44','D2':'#D9812B','D3':'#B54050'}

def save_panel(fig, stem):
    print('save',stem,flush=True)
    for ext in ['png','pdf','svg']:
        fp=OUT/'individual_panels'/f'{stem}.{ext}'
        fig.savefig(fp, dpi=300 if ext=='png' else None, bbox_inches='tight', transparent=True)
    plt.close(fig)
# audit
inventory=[]
for ws in wb.worksheets:
    h=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
    inventory.append({'filename':CNMF.name+'::'+ws.title,'file_type':'xlsx sheet','rows_or_members':'see workbook','columns_or_keys':'; '.join([str(x) for x in h[:8] if x is not None]),'panel_use':'2a/2c/2d/2e/2f','status':'available'})
if CONC.exists():
    conc_tmp=pd.read_csv(CONC,nrows=3)
    inventory.append({'filename':str(CONC.relative_to(BASE)),'file_type':'csv','rows_or_members':pd.read_csv(CONC,usecols=[0]).shape[0],'columns_or_keys':'; '.join(conc_tmp.columns[:8]),'panel_use':'2b','status':'available'})
pd.DataFrame(inventory).to_csv(OUT/'data_audit/file_inventory.tsv',sep='\t',index=False)
# counts
target_names=targets['target_name'].dropna().astype(str).unique()
control=[x for x in target_names if 'non' in x.lower() or 'target' in x.lower()]
counts={'programs':int(summary['program_name'].nunique()),'target_labels':len(target_names),'non_control_target_labels':len(target_names)-len(control),'control_target_labels':len(control),'crt_rows':int(len(assoc))}
json.dump(counts,open(OUT/'cleaned_data/figure2_counts.json','w'),indent=2)
# Fig2a
print('panel2a',flush=True)
fig,ax=plt.subplots(figsize=(6.6,3.2)); ax.set_xlim(0,10); ax.set_ylim(0,5); ax.axis('off')
ax.text(0.05,4.75,'Time-resolved CC-Perturb-seq of endothelial differentiation regulators',fontsize=12,weight='bold')
cards=[(0.45,2.45,2.0,1.2,'Focused library',f"244 screen-derived\n+ curated targets\n{counts['non_control_target_labels']} non-control labels"),(3.05,2.45,2.0,1.2,'CRISPRi hiPSCs','high-MOI guide delivery\n2 biological replicates\ncircular guide capture'),(5.55,2.45,1.8,1.2,'D0-D3','cross-sectional\ncollections\nD0, D1, D2, D3'),(7.85,2.45,1.8,1.2,'Program testing',f"cNMF: {counts['programs']} programs\nper-day CRT\nprogram usage effects")]
for x,y,w,h,title,body in cards:
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.035,rounding_size=0.08',fc='white',ec='#333333',lw=0.9))
    ax.text(x+w/2,y+h-0.25,title,ha='center',va='top',fontsize=9,weight='bold')
    ax.text(x+w/2,y+0.18,body,ha='center',va='bottom',fontsize=7,linespacing=1.25)
for x1,x2 in [(2.45,3.05),(5.05,5.55),(7.35,7.85)]: ax.add_patch(FancyArrowPatch((x1,3.05),(x2,3.05),arrowstyle='-|>',mutation_scale=14,lw=1.2,color='#333333'))
xs=[5.55,6.05,6.55,7.05]
for x,d in zip(xs,['D0','D1','D2','D3']): ax.add_patch(Circle((x,1.35),0.13,fc=DAY[d],ec='white',lw=0.6)); ax.text(x,1.02,d,ha='center',va='top',fontsize=8,weight='bold',color=DAY[d])
ax.plot(xs,[1.35]*4,color='#999999',lw=0.9)
ax.text(0.45,0.45,'Guide-level target origin and exact guide counts require final guide annotation.',fontsize=6.5,color='#666666')
save_panel(fig,'Fig2a_CC_Perturbseq_workflow')
# Fig2b
print('panel2b',flush=True)
conc=pd.read_csv(CONC).dropna(subset=['screen_log2FC','D3_artEC_fraction_shift_pct']).copy()
conc['screen_significant']=conc['screen_fdr']<0.05; conc['perturbseq_significant']=conc['D3_EC_fdr']<0.05
conc['significance_class']=np.select([conc['screen_significant']&conc['perturbseq_significant'], conc['screen_significant'], conc['perturbseq_significant']], ['Both','Screen only','CC-Perturb-seq only'], default='Neither')
labels=['CDH5','KDR','ETV2','ACVR1B','ACVR2A','ACVR2B','BMPR1A','SMAD1','SMAD4','SMAD5','NF2','LATS2','SAV1','TAOK1','AMOTL2']
conc['label_status']=conc['gene_mapped'].isin(labels)
conc.to_csv(OUT/'source_data/Fig2b_screen_Perturbseq_concordance.tsv',sep='\t',index=False)
r,p=pearsonr(conc['screen_log2FC'],conc['D3_artEC_fraction_shift_pct'])
fig,ax=plt.subplots(figsize=(4.25,3.5)); color={'Neither':'#CFCFCF','Screen only':'#E69F00','CC-Perturb-seq only':'#56B4E9','Both':'#7B3294'}
for cl in ['Neither','Screen only','CC-Perturb-seq only','Both']:
    sub=conc[conc['significance_class']==cl]; ax.scatter(sub['screen_log2FC'],sub['D3_artEC_fraction_shift_pct'],s=18 if cl=='Neither' else 30,c=color[cl],edgecolors='white',linewidths=0.25,alpha=0.85,label=cl)
coef=np.polyfit(conc['screen_log2FC'],conc['D3_artEC_fraction_shift_pct'],1); xx=np.linspace(conc['screen_log2FC'].min(),conc['screen_log2FC'].max(),100); ax.plot(xx,coef[0]*xx+coef[1],color='black',lw=1.1)
ax.axhline(0,color='#999999',ls='--',lw=0.7); ax.axvline(0,color='#999999',ls='--',lw=0.7)
texts=[]
for _,row in conc[conc['label_status']].iterrows(): texts.append(ax.text(row['screen_log2FC'],row['D3_artEC_fraction_shift_pct'],row['gene_mapped'],fontsize=6.0))
if adjust_text and texts: adjust_text(texts,ax=ax,arrowprops=dict(arrowstyle='-',lw=0.3,color='#777777'),time_lim=0.8)
ax.text(0.03,0.97,f'n={len(conc)} genes\nr={r:.2f}, P={p:.1e}',transform=ax.transAxes,ha='left',va='top',fontsize=7.4)
ax.set_title('Endpoint concordance',fontsize=9,weight='bold'); ax.set_xlabel('Genome-wide screen effect\n(u-test log2FC)',fontsize=8); ax.set_ylabel('D3 artEC fraction shift\n(target - NT, percentage points)',fontsize=8); ax.legend(frameon=False,fontsize=5.7,loc='lower right'); ax.spines[['top','right']].set_visible(False)
save_panel(fig,'Fig2b_screen_Perturbseq_concordance')
# Fig2c
print('panel2c',flush=True)
usage=[]
for pprog in selected:
    row=summary[summary['program_name']==pprog].iloc[0]
    for d in ['D0','D1','D2','D3']:
        usage.append({'program':pprog,'program_label':f"P{pprog} {prog_labels[pprog][2]}",'dominant_day':prog_labels[pprog][0],'day':d,'mean_usage':float(row[f'Mean program score {d}']),'fraction_above_mean':float(row[f'Fra cells above mean program score {d}']),'top_genes':row['top10_loaded_genes']})
usage=pd.DataFrame(usage); usage['mean_usage_row_z']=usage.groupby('program')['mean_usage'].transform(lambda x:(x-x.mean())/(x.std(ddof=0)+1e-9)); usage.to_csv(OUT/'source_data/Fig2c_selected_program_usage.tsv',sep='\t',index=False)
fig,ax=plt.subplots(figsize=(4.8,4.2)); order=selected[::-1]; py={p:i for i,p in enumerate(order)}; dx={d:i for i,d in enumerate(['D0','D1','D2','D3'])}
sc=None
for _,row in usage.iterrows(): sc=ax.scatter(dx[row['day']],py[row['program']],s=20+row['fraction_above_mean']*220,c=row['mean_usage_row_z'],cmap='viridis',vmin=-1.6,vmax=1.6,marker='s',edgecolors='white',linewidths=0.4)
cb=fig.colorbar(sc,ax=ax,fraction=0.045,pad=0.02); cb.set_label('row-scaled mean\nprogram usage',fontsize=7); cb.ax.tick_params(labelsize=6)
ax.set_xticks(range(4)); ax.set_xticklabels(['D0','D1','D2','D3'],fontsize=8,weight='bold'); ax.set_yticks(range(len(order))); ax.set_yticklabels([f"P{p}  {prog_labels[p][2]}" for p in order],fontsize=6.4); ax.set_xlim(-.5,3.5); ax.set_ylim(-.5,len(order)-.5); ax.grid(color='#EEEEEE',lw=.5); ax.set_title('Selected program usage across D0-D3',fontsize=9,weight='bold'); ax.spines[['top','right']].set_visible(False)
for s,lab in zip([0.15,0.35,0.55],['15%','35%','55%']): ax.scatter([],[],s=20+s*220,c='#808080',marker='s',label=lab)
ax.legend(title='cells above\nmean usage',frameon=False,fontsize=6,title_fontsize=6,loc='lower left',bbox_to_anchor=(1.03,0))
save_panel(fig,'Fig2c_program_usage_atlas')
# helper for assoc features
def effect_matrix(features):
    out=None
    for day,prog,label in features:
        sub=assoc[(assoc['Sample']==day)&(assoc['program_name']==prog)][['target_name','approx_log2FC','adj_pval']].copy(); sub=sub.rename(columns={'approx_log2FC':label,'adj_pval':label+'_q'})
        out=sub if out is None else out.merge(sub,on='target_name',how='inner')
    return out
# Fig2d
print('panel2d',flush=True)
features=[('D1',13,'D1 P13'),('D1',30,'D1 P30'),('D2',15,'D2 P15'),('D2',17,'D2 P17'),('D2',19,'D2 P19'),('D2',23,'D2 P23'),('D3',2,'D3 P2'),('D3',3,'D3 P3'),('D3',11,'D3 P11'),('D3',27,'D3 P27'),('D3',34,'D3 P34')]
mat=effect_matrix(features)
groups={'TGF/BMP/Activin':['ACVR1B','ACVR2A','ACVR2B','BMPR1A','SMAD1','SMAD3','SMAD4','SMAD5','TGFBR2','LEMD3'],'VEGF/EC specification':['KDR','ETV2','PLCG1','CDH5','NOTCH1','KCTD10','HAND1'],'Hippo/YAP':['NF2','LATS2','SAV1','TAOK1','AMOTL2','WWC3'],'Chromatin/RNA':['SMARCA4','SETDB1','SUZ12','MED12','ZNF281','EOMES','MIXL1','TBXT','TBX3'],'Cell fitness/process':['VHL','BAX','CASP3','RACK1','TOP1','EIF4E']}
rows=[]
for g,genes in groups.items():
    for gene in genes:
        if gene in set(mat['target_name']): rows.append({'target_name':gene,'group':g})
regsel=pd.DataFrame(rows); heat=regsel.merge(mat,on='target_name',how='left'); heat.to_csv(OUT/'source_data/Fig2d_regulator_program_map.tsv',sep='\t',index=False)
vals=heat[[f[2] for f in features]].astype(float).values; clip=max(1.0,float(np.nanpercentile(np.abs(vals),98))); clip=round(clip,2)
fig,ax=plt.subplots(figsize=(5.6,5.1)); im=ax.imshow(np.clip(vals,-clip,clip),aspect='auto',cmap='RdBu_r',norm=TwoSlopeNorm(vmin=-clip,vcenter=0,vmax=clip))
for i,row in heat.iterrows():
    for j,f in enumerate(features):
        q=row[f[2]+'_q']
        if pd.notna(q) and q<0.05: ax.plot(j,i,'o',ms=1.8,mfc='black',mec='black')
ax.set_xticks(range(len(features))); ax.set_xticklabels([f[2] for f in features],rotation=45,ha='right',fontsize=6.3); ax.set_yticks(range(len(heat))); ax.set_yticklabels(heat['target_name'],fontsize=5.8)
start=0
for g,sub in heat.groupby('group',sort=False):
    end=start+len(sub); ax.axhline(end-0.5,color='white',lw=1.2); ax.text(-1.0,(start+end-1)/2,g,ha='right',va='center',fontsize=6.1); start=end
cb=fig.colorbar(im,ax=ax,fraction=.04,pad=.02); cb.set_label(f'CRT effect on\nprogram usage\nclipped at ±{clip:.2f}',fontsize=6.5); cb.ax.tick_params(labelsize=6)
ax.set_title('Regulator effects on selected programs',fontsize=9,weight='bold'); ax.tick_params(length=0); [sp.set_visible(False) for sp in ax.spines.values()]
pd.DataFrame([{'display_clip_abs':clip,'original_min':float(np.nanmin(vals)),'original_max':float(np.nanmax(vals)),'n_clipped':int((np.abs(vals)>clip).sum())}]).to_csv(OUT/'source_data/Fig2d_display_scale.tsv',sep='\t',index=False)
save_panel(fig,'Fig2d_regulator_program_map')
# Fig2e
print('panel2e',flush=True)
net_features=[('D1',30,'P30'),('D1',13,'P13'),('D2',15,'P15'),('D2',17,'P17'),('D3',3,'P3'),('D3',34,'P34')]
net=effect_matrix(net_features); pairs=[('P30','P15'),('P30','P17'),('P13','P15'),('P13','P17'),('P15','P17'),('P15','P3'),('P15','P34'),('P17','P3'),('P17','P34'),('P3','P34')]; pr=[]
for a,b in pairs:
    valid=(net['target_name']!='non-targeting')&net[a].notna()&net[b].notna(); rr,pp=pearsonr(net.loc[valid,a].astype(float),net.loc[valid,b].astype(float)); pr.append({'source':a,'target':b,'pearson_r':rr,'p_value':pp,'n_regulators':int(valid.sum())})
pr=pd.DataFrame(pr); pr.to_csv(OUT/'source_data/Fig2e_pairwise_correlations.tsv',sep='\t',index=False)
pos={'P30':(0,1),'P13':(0,-1),'P15':(2.1,1),'P17':(2.1,-1),'P3':(4.2,1),'P34':(4.2,-1)}; labels={'P30':'P30\nanterior PS','P13':'P13\nposterior PS','P15':'P15\nPre-EC','P17':'P17\nLat-Meso','P3':'P3\nEC endpoint','P34':'P34\ncontractile alt.'}
fig,ax=plt.subplots(figsize=(5.6,3.6)); ax.axis('off'); ax.set_xlim(-.8,5); ax.set_ylim(-1.9,1.8)
for x,d in [(0,'D1'),(2.1,'D2'),(4.2,'D3')]: ax.text(x,1.58,d,ha='center',va='center',fontsize=11,weight='bold')
edge_set=[('P30','P15'),('P30','P17'),('P13','P15'),('P13','P17'),('P15','P17'),('P15','P3'),('P15','P34'),('P17','P3'),('P17','P34')]
for _,rrr in pr[pr.apply(lambda r:(r['source'],r['target']) in edge_set,axis=1)].iterrows():
    a,b=rrr['source'],rrr['target']; x1,y1=pos[a]; x2,y2=pos[b]; col=POS if rrr['pearson_r']>0 else NEG; lw=.7+3.0*abs(rrr['pearson_r']); style='-' if abs(rrr['pearson_r'])>=.30 else '--'
    if set([a,b])==set(['P15','P17']): ax.add_patch(FancyArrowPatch((x1,y1-.35),(x2,y2+.35),arrowstyle='<->',mutation_scale=10,lw=lw,color=col,linestyle=style)); ax.text(x1+.35,0,f"r={rrr['pearson_r']:.2f}",color=col,fontsize=7,weight='bold')
    else:
        ax.add_patch(FancyArrowPatch((x1+.38,y1),(x2-.38,y2),arrowstyle='-|>',mutation_scale=11,lw=lw,color=col,linestyle=style,alpha=.95)); ax.text((x1+x2)/2,(y1+y2)/2+(0.13 if y1==y2 else 0),f"{rrr['pearson_r']:+.2f}",color=col,fontsize=7,weight='bold',ha='center')
for n,(x,y) in pos.items(): ax.add_patch(FancyBboxPatch((x-.42,y-.28),.84,.56,boxstyle='round,pad=.03,rounding_size=.08',fc='#FFFFFF',ec='#333333',lw=.8,zorder=3)); ax.text(x,y,labels[n],ha='center',va='center',fontsize=7.2,zorder=4)
ax.text(.05,-1.65,'Edges are correlations of perturbation effects across separate samples; not lineage tracing.',fontsize=6.2,color='#555555'); ax.set_title('Cross-time perturbation-effect relationships',fontsize=9,weight='bold')
save_panel(fig,'Fig2e_cross_time_relationships')
# Fig2f
print('panel2f',flush=True)
sel6=[30,13,15,17,3,34]; markers=[]
for p in sel6:
    sub=loadings[loadings['Program'].astype('Int64')==p].sort_values('Rank').head(4)
    for _,r in sub.iterrows(): markers.append({'program':p,'gene':r['Gene'],'rank_within_program':r['Rank'],'selection_reason':'top 4 loading rank'})
markers=pd.DataFrame(markers); genes=markers['gene'].unique().tolist(); matrix=[]
for p in sel6:
    ranks=dict(zip(loadings[(loadings['Program'].astype('Int64')==p)&(loadings['Gene'].isin(genes))]['Gene'], loadings[(loadings['Program'].astype('Int64')==p)&(loadings['Gene'].isin(genes))]['Rank']))
    for g in genes:
        rank=ranks.get(g,np.nan); score=(31-rank)/30 if pd.notna(rank) and rank<=30 else 0; matrix.append({'program':p,'gene':g,'rank_within_program':rank,'rank_score':score})
mr=pd.DataFrame(matrix); markers.to_csv(OUT/'source_data/Fig2f_marker_selection.tsv',sep='\t',index=False); mr.to_csv(OUT/'source_data/Fig2f_marker_rank_matrix.tsv',sep='\t',index=False)
reg_ann=[]; prog_day={30:'D1',13:'D1',15:'D2',17:'D2',3:'D3',34:'D3'}
for p in sel6:
    sub=assoc[(assoc['Sample']==prog_day[p])&(assoc['program_name']==p)&(assoc['target_name']!='non-targeting')]; sig=sub[sub['adj_pval']<0.05]; sig=sig if len(sig)>0 else sub
    reg_ann.append({'program':p,'positive':', '.join(sig.sort_values('approx_log2FC',ascending=False).head(2)['target_name']),'negative':', '.join(sig.sort_values('approx_log2FC').head(2)['target_name'])})
reg_ann=pd.DataFrame(reg_ann); reg_ann.to_csv(OUT/'source_data/Fig2f_regulator_effect_source_data.tsv',sep='\t',index=False)
fig=plt.figure(figsize=(6.3,4.6)); gs=fig.add_gridspec(1,2,width_ratios=[1.8,1.2],wspace=.18); ax=fig.add_subplot(gs[0,0]); py={p:i for i,p in enumerate(sel6[::-1])}; gx={g:i for i,g in enumerate(genes)}; arr=np.zeros((len(sel6),len(genes)))
for _,row in mr.iterrows(): arr[py[row['program']],gx[row['gene']]]=row['rank_score']
im=ax.imshow(arr,aspect='auto',cmap='YlGnBu',vmin=0,vmax=1); ax.set_xticks(range(len(genes))); ax.set_xticklabels(genes,rotation=60,ha='right',fontsize=5.3); ax.set_yticks(range(len(sel6))); ax.set_yticklabels([f"P{p} {prog_labels[p][2]}" for p in sel6[::-1]],fontsize=6.3); ax.set_title('Top-loading marker genes',fontsize=9,weight='bold'); ax.tick_params(length=0); [sp.set_visible(False) for sp in ax.spines.values()]; cb=fig.colorbar(im,ax=ax,fraction=.04,pad=.02); cb.set_label('rank score\n(top 30)',fontsize=6); cb.ax.tick_params(labelsize=5.8)
ax2=fig.add_subplot(gs[0,1]); ax2.axis('off'); ax2.set_title('Representative regulator effects',fontsize=9,weight='bold',loc='left'); y=.95
for p in sel6:
    rr=reg_ann[reg_ann['program']==p].iloc[0]; ax2.text(0,y,f"P{p}  {prog_labels[p][2]}",fontsize=7.0,weight='bold',ha='left',va='top'); ax2.text(.02,y-.07,'+ '+rr['positive'],fontsize=6.0,color=POS,ha='left',va='top'); ax2.text(.02,y-.13,'- '+rr['negative'],fontsize=6.0,color=NEG,ha='left',va='top'); y-=.16
ax2.text(0,-.03,'Motif/logo input unavailable; marker matrix uses program-loading rank.',fontsize=6.0,color='#555555')
save_panel(fig,'Fig2f_selected_program_annotation')
# assemble
print('assemble',flush=True)
fig=plt.figure(figsize=(13.5,10)); gs=fig.add_gridspec(3,2,height_ratios=[.78,1.1,1.05],width_ratios=[1,1.05],hspace=.24,wspace=.16); panel={'a':'Fig2a_CC_Perturbseq_workflow.png','b':'Fig2b_screen_Perturbseq_concordance.png','c':'Fig2c_program_usage_atlas.png','d':'Fig2d_regulator_program_map.png','e':'Fig2e_cross_time_relationships.png','f':'Fig2f_selected_program_annotation.png'}; pos={'a':gs[0,0],'b':gs[0,1],'c':gs[1,0],'d':gs[1,1],'e':gs[2,0],'f':gs[2,1]}
for letter in 'abcdef':
    ax=fig.add_subplot(pos[letter]); ax.axis('off'); img=plt.imread(OUT/'individual_panels'/panel[letter]); ax.imshow(img); ax.text(-.02,1.04,letter,transform=ax.transAxes,fontsize=16,weight='bold')
fig.suptitle('Time-resolved CC-Perturb-seq maps regulators to developmental programs',fontsize=16,weight='bold',y=.995)
for ext in ['png','pdf','svg']:
    fig.savefig(OUT/'assembled_figure'/f'Figure2_balanced.{ext}',dpi=300 if ext=='png' else None,bbox_inches='tight')
plt.close(fig)
# docs
print('docs',flush=True)
pd.DataFrame([{'panel':'2a','source_file':CNMF.name,'source_object':'Targets Summary, Summary','use':'counts and workflow'}, {'panel':'2b','source_file':str(CONC.relative_to(BASE)),'source_object':'merged table','use':'concordance'}, {'panel':'2c','source_file':CNMF.name,'source_object':'Summary','use':'program usage'}, {'panel':'2d','source_file':CNMF.name,'source_object':'Perturbation Association 1','use':'CRT effects'}, {'panel':'2e','source_file':CNMF.name,'source_object':'Perturbation Association 1','use':'effect correlations'}, {'panel':'2f','source_file':CNMF.name,'source_object':'Program Loadings + Perturbation Association 1','use':'markers and regulators'}]).to_csv(OUT/'data_audit/file_to_panel_map.tsv',sep='\t',index=False)
open(OUT/'data_audit/unresolved_issues.md','w').write('# Unresolved Figure 2 input issues\n\n- Cell metadata/UMAP coordinates were not available; panel 2c uses cNMF Summary program usage rather than UMAPs.\n- Final guide annotation was not available; panel 2a reports cNMF target labels, not guide counts.\n- Motif matrices/logos were not available; panel 2f excludes motifs.\n- Cross-time edges are correlations across perturbation effects in separate samples, not lineage tracing.\n')
open(OUT/'README_figure2_build.md','w').write(f'# Figure 2 build\n\nGenerated from `{CNMF.name}` and available screen-to-D3 concordance source.\n\nCounts: {counts}\n')
open(OUT/'legends/Figure2_legend_draft.md','w').write(f'# Figure 2 legend draft\n\n**a,** CC-Perturb-seq workflow and analysis design. **b,** Concordance between genome-wide screen u-test log2FC and D3 artEC fraction shift for {len(conc)} matched genes. **c,** Selected cNMF program usage across D0-D3. **d,** Per-day CRT regulator effects on program usage. **e,** Cross-time perturbation-effect correlations; these are not lineage tracing. **f,** Selected program annotation using top-loading genes and representative regulators; motifs were not available.\n')
open(OUT/'final_QC_checklist.md','w').write('# QC checklist\n\n- Six panels exported as SVG/PDF/PNG.\n- Assembled Figure 2 exported as SVG/PDF/PNG.\n- Missing UMAP/guide/motif inputs documented.\n')
shutil.copy(__file__, OUT/'scripts/build_figure2_simple.py')
zip_path=BASE/'figure2_build_simple_package.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
    for root,dirs,files in os.walk(OUT):
        for fn in files:
            fp=Path(root)/fn; z.write(fp,arcname=str(fp.relative_to(OUT.parent)))
print('DONE',zip_path,flush=True)
